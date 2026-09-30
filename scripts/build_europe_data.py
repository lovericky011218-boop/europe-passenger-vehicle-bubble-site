from __future__ import annotations

import json
import math
import re
import statistics
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CURRENT = Path('/Users/ricky/Desktop/海外V/外派/总体数据/欧洲1-7月.xlsx')
GLOBAL_DATA = Path('/Users/ricky/Desktop/海外V/欧洲&南美调研/调研方案/global-vehicle-bubble-site/dist/data.js')
MANUAL = ROOT / 'data/manual-europe-prices.json'
OUT = ROOT / 'dist/data.js'
AUDIT = ROOT / 'data/price-audit.json'
SPEC_OVERRIDES = ROOT / 'data/manual-europe-specs.json'
IMPORT_AUDIT = ROOT / 'data/import-audit.json'

ALLOWED_BODIES = {'Car', 'SUV', 'MPV'}
ALLOWED_FUELS = {'ICE', 'HEV', 'MEV', 'REEV', 'BEV', 'PHEV'}
COUNTRY_ALIASES = {'Russia': 'Russia (Sales)'}


def clean(value):
    return str(value or '').strip()


def finite(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def med(values):
    values = [float(value) for value in values if finite(value) is not None]
    return statistics.median(values) if values else None


def norm(value):
    text = unicodedata.normalize('NFKD', clean(value)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def slug(value):
    return re.sub(r'[^a-z0-9]+', '-', norm(value)).strip('-') or 'vehicle'


def parse_global():
    text = GLOBAL_DATA.read_text('utf-8')
    start = text.index('window.VEHICLE_DATA=') + len('window.VEHICLE_DATA=')
    end = text.index(';\nwindow.REFERENCE_DATA=')
    return json.loads(text[start:end])


global_rows = [row for row in parse_global() if row.get('continent') == 'Europe']
manual_rows = json.loads(MANUAL.read_text('utf-8'))
spec_overrides = json.loads(SPEC_OVERRIDES.read_text('utf-8'))

exact = defaultdict(list)
same_fuel = defaultdict(list)
same_model = defaultdict(list)
for row in global_rows:
    # The previous import already converts price to EUR; localPrice is NOT EUR.
    price = finite(row.get('price'))
    if not price or price <= 0:
        continue
    country = clean(row.get('country'))
    model_key = norm(row.get('model'))
    fuel = clean(row.get('fuel'))
    exact[(country, model_key, fuel)].append(row)
    same_fuel[(model_key, fuel)].append(row)
    same_model[model_key].append(row)

manual = {(row.get('country'), norm(row['model']), row['fuel']): row for row in manual_rows}
country_rates = defaultdict(list)
for row in global_rows:
    eur, local = finite(row.get('price')), finite(row.get('localPrice'))
    if eur and eur > 0 and local and local > 0:
        country_rates[row['country']].append(local / eur)


def source_price(rows, method, year):
    preferred = [row for row in rows if row.get('year') == year] or rows
    prices = [finite(row.get('price')) for row in preferred]
    mins = [finite(row.get('priceMin')) or finite(row.get('price')) for row in preferred]
    maxs = [finite(row.get('priceMax')) or finite(row.get('price')) for row in preferred]
    return {
        'price': round(med(prices)),
        'priceMin': round(min(value for value in mins if value is not None)),
        'priceMax': round(max(value for value in maxs if value is not None)),
        'priceMethod': method,
        'priceSource': '欧洲合并1-4.xlsx',
        'priceNote': 'Matched from Price. (EUR) in prior Europe workbook-derived price map',
    }


def resolve_price(country, model, fuel, year):
    key = (norm(model), fuel)
    old_country = COUNTRY_ALIASES.get(country, country)
    manual_row = manual.get((country, *key)) or manual.get((None, *key))
    if manual_row and manual_row.get('localPrice'):
        rate = med(country_rates[old_country])
        if not rate:
            raise ValueError(f'No workbook exchange-rate basis for {country}')
        value = round(manual_row['localPrice'] / rate)
        return {
            'price': value, 'priceMin': value, 'priceMax': value,
            'priceMethod': 'web_public_median_converted', 'priceSource': manual_row['source'],
            'priceNote': f"{manual_row.get('note', '')}; fixed workbook rate {rate:.6f} local units/EUR",
        }
    if exact.get((old_country, key[0], fuel)):
        return source_price(exact[(old_country, key[0], fuel)], 'country_exact', year)
    if same_fuel.get(key):
        return source_price(same_fuel[key], 'europe_model_fuel_median', year)
    if manual_row:
        row = manual_row
        return {
            'price': round(row['price']), 'priceMin': round(row.get('priceMin', row['price'])),
            'priceMax': round(row.get('priceMax', row['price'])), 'priceMethod': 'web_public_median',
            'priceSource': row['source'], 'priceNote': row.get('note', ''),
        }
    if same_model.get(key[0]):
        result = source_price(same_model[key[0]], 'europe_model_median_other_powertrain', year)
        result['priceNote'] = 'Model matched in prior Europe workbook; powertrain-specific price unavailable'
        return result
    return {'price': None, 'priceMin': None, 'priceMax': None, 'priceMethod': 'unavailable', 'priceSource': '', 'priceNote': 'No verified public MSRP found'}


wb = load_workbook(CURRENT, read_only=True, data_only=True)
ws = wb.active
headers = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
indices = {name: index for index, name in enumerate(headers)}
groups = {}
known_lengths = defaultdict(list)
for row in global_rows:
    if finite(row.get('length')) and row['length'] > 0:
        known_lengths[norm(row['model'])].append(row['length'])
source_totals = defaultdict(float)
excluded = defaultdict(int)
for values in ws.iter_rows(min_row=2, values_only=True):
    country = clean(values[indices['Country/Territory-Name']])
    body = clean(values[indices['Global Sales Sub-Segment']])
    model = clean(values[indices['Model (World)']])
    fuel = clean(values[indices['Fuel Type']])
    length = finite(values[indices['Length']])
    wheelbase = finite(values[indices['Wheelbase']])
    if not country or body not in ALLOWED_BODIES or fuel not in ALLOWED_FUELS or not model or 'unspec' in norm(model):
        excluded['outside_passenger_scope_or_unspecified'] += 1
        continue
    if length and length > 0:
        known_lengths[norm(model)].append(length)
    for year in (2024, 2025, 2026):
        sales = finite(values[indices[f'CY{year}']])
        if not sales or sales <= 0:
            continue
        key = (country, year, model, fuel, body)
        group = groups.setdefault(key, {'sales': 0, 'lengths': [], 'wheelbases': [], 'sourceRows': 0})
        group['sales'] += sales
        source_totals[(country, year)] += sales
        if length and length > 0:
            group['lengths'].append(length)
        if wheelbase and wheelbase > 0:
            group['wheelbases'].append(wheelbase)
        group['sourceRows'] += 1
wb.close()

rows = []
method_counts = defaultdict(int)
method_sales = defaultdict(float)
length_methods = defaultdict(int)
countries = set()
for (country, year, model, fuel, body), group in groups.items():
    price = resolve_price(country, model, fuel, year)
    length = med(group['lengths'])
    length_method = 'source_workbook'
    length_source = CURRENT.name
    if length is None:
        override = spec_overrides.get(model)
        if override:
            length, length_method, length_source = override['length'], 'official_web', override['source']
        else:
            length = med(known_lengths[norm(model)])
            length_method = 'same_model_median' if length else 'unavailable'
            length_source = '欧洲源表 / 欧洲合并1-4.xlsx' if length else ''
    length_methods[length_method] += 1
    sales = round(group['sales'])
    method_counts[price['priceMethod']] += 1
    method_sales[price['priceMethod']] += sales
    countries.add(country)
    rows.append({
        'id': f"{slug(country)}-{year}-{slug(model)}-{fuel.lower()}-{body.lower()}",
        'country': country, 'year': year,
        'brand': model.split()[0], 'model': model, 'fuel': fuel, 'body': body,
        'length': round(length) if length else None, 'lengthMethod': length_method,
        'lengthSource': length_source,
        'price': price['price'], 'priceMin': price['priceMin'], 'priceMax': price['priceMax'],
        'priceMethod': price['priceMethod'], 'priceSource': price['priceSource'],
        'priceNote': price['priceNote'], 'sales': sales, 'sourceRows': group['sourceRows'],
    })

rows.sort(key=lambda row: (row['country'], row['year'], row['model'], row['fuel'], row['body']))
imported_totals = defaultdict(int)
for row in rows:
    imported_totals[(row['country'], row['year'])] += row['sales']
assert dict(imported_totals) == {key: round(value) for key, value in source_totals.items()}, 'Import lost or duplicated source sales'
coverage = [
    {'country': country, 'year': year, 'sourceSales': round(value), 'importedSales': imported_totals[(country, year)]}
    for (country, year), value in sorted(source_totals.items())
]
incomplete = [row for row in rows if row['price'] is None or row['length'] is None]
IMPORT_AUDIT.write_text(json.dumps({'coverage': coverage, 'excludedRows': dict(excluded), 'lengthMethods': dict(length_methods), 'incompleteRows': incomplete}, ensure_ascii=False, indent=2), encoding='utf-8')
refs = [
    {'id': 'reference-g01', 'model': 'G01', 'length': 4500, 'price': 30000, 'body': 'SUV', 'reference': True},
    {'id': 'reference-g02', 'model': 'G02', 'length': 4650, 'price': 35000, 'body': 'SUV', 'reference': True},
]
meta = {
    'years': [2024, 2025, 2026], 'latestYear': 2026, 'latestPeriod': '2026 YTD (Jan–Jul)',
    'rows': len(rows), 'countries': sorted(countries), 'currency': 'EUR',
    'sourceWorkbook': '欧洲1-7月.xlsx', 'priceWorkbook': '欧洲合并1-4.xlsx',
    'method': 'Passenger Car/SUV/MPV only; Unspecified removed; all positive source sales retained, including missing dimensions/prices; sales summed by country + year + model + fuel + body.',
    'priceCurrencyBasis': 'Price. (EUR) from Europe workbook; never localPrice without conversion',
    'lengthMethods': dict(length_methods), 'incompleteRows': len(incomplete),
    'priceMethods': dict(sorted(method_counts.items())),
    'priceMethodSales': {key: round(value) for key, value in sorted(method_sales.items())},
    'webSources': manual_rows,
}
OUT.write_text(
    'window.VEHICLE_DATA=' + json.dumps(rows, ensure_ascii=False, separators=(',', ':')) + ';\n'
    + 'window.REFERENCE_DATA=' + json.dumps(refs, ensure_ascii=False, separators=(',', ':')) + ';\n'
    + 'window.DATA_META=' + json.dumps(meta, ensure_ascii=False, separators=(',', ':')) + ';\n',
    encoding='utf-8',
)
AUDIT.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'rows': len(rows), 'countries': len(countries), 'priceMethods': meta['priceMethods'], 'priceMethodSales': meta['priceMethodSales']}, ensure_ascii=False, indent=2))
